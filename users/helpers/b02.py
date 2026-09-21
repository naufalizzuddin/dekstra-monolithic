from copy import copy
from datetime import datetime
import io
import os
import tempfile

import qrcode
from openpyxl import load_workbook
from openpyxl.styles import Font, Alignment
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string


SHEET_NAME = "F1.02"
QR_CELL = "L43"


# =========================================================
# STYLE
# =========================================================
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
CHECK_FONT = Font(name="Arial", size=11, bold=True)
CHECK_MARK = "✓"


# =========================================================
# HELPER DASAR
# =========================================================
def normalize_text(value):
    if value is None:
        return ""
    return str(value).strip()


def set_value(ws, cell_ref, value, align=None):
    if value is None or value == "":
        return
    ws[cell_ref] = value
    if align:
        ws[cell_ref].alignment = copy(align)


def set_merged_value(ws, top_left_cell, value, align=None):
    if value is None or value == "":
        return
    ws[top_left_cell] = value
    if align:
        ws[top_left_cell].alignment = copy(align)


def mark_checkbox(ws, cell_ref, mark=CHECK_MARK):
    ws[cell_ref] = mark
    ws[cell_ref].alignment = copy(CENTER)
    ws[cell_ref].font = copy(CHECK_FONT)


def normalize_date(value):
    value = normalize_text(value)
    if not value:
        return None

    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%Y/%m/%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(value, fmt)
        except ValueError:
            continue
    return None


def build_footer_date_text(raw_data):
    tempat = normalize_text(raw_data.get("tempat", ""))
    tanggal = normalize_text(raw_data.get("tanggal", ""))
    bulan = normalize_text(raw_data.get("bulan", ""))
    tahun = normalize_text(raw_data.get("tahun", ""))

    # Prioritas: jika models.py mengirim tanggal_ttd format YYYY-MM-DD
    parsed = normalize_date(raw_data.get("tanggal_ttd", ""))
    if parsed:
        tanggal = str(parsed.day)
        bulan = parsed.strftime("%B")
        tahun = str(parsed.year)

    if tempat or tanggal or bulan or tahun:
        return f"{tempat}, {tanggal} {bulan} {tahun}".strip().strip(",")

    return ""


def create_qr_image_file(text):
    if not text:
        return None

    temp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".png")
    temp_file.close()

    qr = qrcode.QRCode(
        version=2,
        error_correction=qrcode.constants.ERROR_CORRECT_M,
        box_size=8,
        border=1,
    )
    qr.add_data(text)
    qr.make(fit=True)

    img = qr.make_image(fill_color="black", back_color="white")
    img.save(temp_file.name)
    return temp_file.name


def safe_remove_file(path):
    try:
        if path and os.path.exists(path):
            os.remove(path)
    except Exception:
        pass


def insert_qr_image(ws, image_path, cell_ref, width=50, height=50):
    if not image_path or not os.path.exists(image_path):
        return

    img = XLImage(image_path)
    img.width = width
    img.height = height

    col_letter, row_number = coordinate_from_string(cell_ref)
    col_idx = column_index_from_string(col_letter) - 1
    row_idx = row_number - 1

    # Ukuran default Excel jika belum diset
    col_width = ws.column_dimensions[col_letter].width or 8.43
    row_height = ws.row_dimensions[row_number].height or 15

    # Konversi kasar Excel unit ke pixel
    cell_width_px = int(col_width * 7)
    cell_height_px = int(row_height * 1.33)

    # Offset agar gambar berada di tengah cell
    offset_x_px = max((cell_width_px - width) // 2, 0)
    offset_y_px = max((cell_height_px - height) // 2, 0)

    marker = AnchorMarker(
        col=col_idx,
        colOff=offset_x_px * 9525,
        row=row_idx,
        rowOff=offset_y_px * 9525,
    )

    size = XDRPositiveSize2D(
        cx=width * 9525,
        cy=height * 9525,
    )

    img.anchor = OneCellAnchor(
        _from=marker,
        ext=size,
    )

    ws.add_image(img)


# =========================================================
# MAPPING DATA PEMOHON
# =========================================================
PEMOHON_CELLS = {
    "nama_lengkap": "F5",
    "nik": "F6",
    "nomor_kk": "F7",
    "nomor_hp_wa": "F8",
}


# =========================================================
# MAPPING JENIS PERMOHONAN
# =========================================================
JENIS_KK_MAP = {
    "kk-baru-membentuk-keluarga": "E15",
    "kk-baru-pergantian-kepala": "E16",
    "kk-baru-pisah-kk": "E17",
    "kk-baru-pindah-datang": "E18",
    "kk-baru-wni-dari-ln": "E19",
    "kk-baru-rentan-adminduk": "E20",
    "kk-ubah-menumpang": "E22",
    "kk-ubah-peristiwa-penting": "E23",
    "kk-ubah-elemen-data": "E24",
    "kk-hilang": "E26",
    "kk-rusak": "E27",
}

JENIS_KTP_MAP = {
    "ktp-baru": "I14",
    "ktp-pindah-datang": "I15",
    "ktp-hilang": "I17",
    "ktp-rusak": "I18",
    "ktp-perpanjangan-itap": "I19",
    "ktp-perubahan-status-kewarganegaraan": "I20",
    "ktp-luar-domisili": "I22",
    "ktp-transmigrasi": "I23",
}

JENIS_KIA_MAP = {
    "kia-baru": "M14",
    "kia-hilang": "M16",
    "kia-rusak": "M17",
    "kia-perpanjangan-itap": "M18",
    "kia-lainnya": "M19",
}

JENIS_PERUBAHAN_DATA_MAP = {
    "ubah-kk": "P14",
    "ubah-ktp": "P15",
    "ubah-kia": "P16",
}


# =========================================================
# MAPPING LAMPIRAN
# =========================================================
LAMPIRAN_MAP = {
    "lampiran_kk_lama": "C30",
    "lampiran_buku_nikah": "C31",
    "lampiran_akta_perceraian": "C32",
    "lampiran_surat_pindah": "C33",
    "lampiran_surat_pindah_luar_negeri": "C34",
    "lampiran_ktp_rusak": "C35",
    "lampiran_dokumen_perjalanan": "C36",
    "lampiran_surat_keterangan_hilang": "C37",
    "lampiran_surat_keterangan_perubahan": "I30",
    "lampiran_sptjm": "I32",
    "lampiran_akta_kematian": "I33",
    "lampiran_surat_pernyataan_hilang_rusak": "I34",
    "lampiran_surat_pindah_perwakilan_ri": "I35",
    "lampiran_surat_pernyataan_anggota": "I36",
    "lampiran_surat_kuasa_pengasuhan": "I37",
    "lampiran_kartu_izin_tinggal_tetap": "I38",
}


# =========================================================
# FOOTER / TANDA TANGAN
# =========================================================
FOOTER_CELLS = {
    "tanggal_lokasi": "K41",
    "nama_pemohon": "D46",
    "nama_petugas": "L46",
}


# =========================================================
# TRANSFORM RAW DATA
# =========================================================
def transform_raw_to_data(raw_data):
    return {
        # Data pemohon
        "nama_lengkap": normalize_text(raw_data.get("nama_lengkap", "")),
        "nik": normalize_text(raw_data.get("nik", "")),
        "nomor_kk": normalize_text(raw_data.get("nomor_kk", "")),
        "nomor_hp_wa": normalize_text(raw_data.get("nomor_hp_wa", "")),

        # Jenis permohonan
        "kategori_permohonan": normalize_text(raw_data.get("kategori_permohonan", "")),
        "jenis_kk": raw_data.get("jenis_kk"),
        "jenis_ktp": raw_data.get("jenis_ktp"),
        "jenis_kia": raw_data.get("jenis_kia"),
        "jenis_perubahan_data": raw_data.get("jenis_perubahan_data"),

        # Lampiran
        "lampiran_kk_lama": bool(raw_data.get("lampiran_kk_lama", False)),
        "lampiran_buku_nikah": bool(raw_data.get("lampiran_buku_nikah", False)),
        "lampiran_akta_perceraian": bool(raw_data.get("lampiran_akta_perceraian", False)),
        "lampiran_surat_pindah": bool(raw_data.get("lampiran_surat_pindah", False)),
        "lampiran_surat_pindah_luar_negeri": bool(raw_data.get("lampiran_surat_pindah_luar_negeri", False)),
        "lampiran_ktp_rusak": bool(raw_data.get("lampiran_ktp_rusak", False)),
        "lampiran_dokumen_perjalanan": bool(raw_data.get("lampiran_dokumen_perjalanan", False)),
        "lampiran_surat_keterangan_hilang": bool(raw_data.get("lampiran_surat_keterangan_hilang", False)),
        "lampiran_surat_keterangan_perubahan": bool(raw_data.get("lampiran_surat_keterangan_perubahan", False)),
        "lampiran_sptjm": bool(raw_data.get("lampiran_sptjm", False)),
        "lampiran_akta_kematian": bool(raw_data.get("lampiran_akta_kematian", False)),
        "lampiran_surat_pernyataan_hilang_rusak": bool(raw_data.get("lampiran_surat_pernyataan_hilang_rusak", False)),
        "lampiran_surat_pindah_perwakilan_ri": bool(raw_data.get("lampiran_surat_pindah_perwakilan_ri", False)),
        "lampiran_surat_pernyataan_anggota": bool(raw_data.get("lampiran_surat_pernyataan_anggota", False)),
        "lampiran_surat_kuasa_pengasuhan": bool(raw_data.get("lampiran_surat_kuasa_pengasuhan", False)),
        "lampiran_kartu_izin_tinggal_tetap": bool(raw_data.get("lampiran_kartu_izin_tinggal_tetap", False)),

        # Footer
        "tanggal_lokasi": build_footer_date_text(raw_data),
        "nama_pemohon": normalize_text(
            raw_data.get("nama_pemohon") or raw_data.get("nama_lengkap", "")
        ),
        "nama_petugas": normalize_text(
            raw_data.get("nama_petugas", "ADMIN")
        ),

        # QR
        "qr_text": normalize_text(raw_data.get("qr_ttd", "")),
    }


# =========================================================
# LOGIC FILL
# =========================================================
def fill_data_pemohon(ws, data):
    for field_name, cell_ref in PEMOHON_CELLS.items():
        set_merged_value(ws, cell_ref, data.get(field_name, ""), LEFT)


def fill_jenis_permohonan(ws, data):
    kategori = data.get("kategori_permohonan")

    if kategori == "kartu-keluarga" and data.get("jenis_kk"):
        target = JENIS_KK_MAP.get(data["jenis_kk"])
        if target:
            mark_checkbox(ws, target)

    elif kategori == "ktp-el" and data.get("jenis_ktp"):
        target = JENIS_KTP_MAP.get(data["jenis_ktp"])
        if target:
            mark_checkbox(ws, target)

    elif kategori == "kia" and data.get("jenis_kia"):
        target = JENIS_KIA_MAP.get(data["jenis_kia"])
        if target:
            mark_checkbox(ws, target)

    elif kategori == "perubahan-data" and data.get("jenis_perubahan_data"):
        target = JENIS_PERUBAHAN_DATA_MAP.get(data["jenis_perubahan_data"])
        if target:
            mark_checkbox(ws, target)


def fill_lampiran(ws, data):
    for field_name, cell_ref in LAMPIRAN_MAP.items():
        if data.get(field_name, False):
            mark_checkbox(ws, cell_ref)


def fill_footer(ws, data):
    set_value(ws, FOOTER_CELLS["tanggal_lokasi"], data.get("tanggal_lokasi", ""), CENTER)
    set_value(ws, FOOTER_CELLS["nama_pemohon"], data.get("nama_pemohon", ""), CENTER)
    set_value(ws, FOOTER_CELLS["nama_petugas"], data.get("nama_petugas", "ADMIN"), CENTER)

    qr_path = create_qr_image_file(data.get("qr_text", ""))
    if qr_path:
        insert_qr_image(ws, qr_path, QR_CELL, width=50, height=50)

    return qr_path


# =========================================================
# MAIN GENERATOR
# =========================================================
def generate_b02_excel(template_path, raw_data):
    data = transform_raw_to_data(raw_data)

    wb = load_workbook(template_path)
    ws = wb[SHEET_NAME]

    fill_data_pemohon(ws, data)
    fill_jenis_permohonan(ws, data)
    fill_lampiran(ws, data)

    qr_path = fill_footer(ws, data)

    output = io.BytesIO()
    try:
        wb.save(output)
        output.seek(0)
        return output.getvalue()
    finally:
        safe_remove_file(qr_path)