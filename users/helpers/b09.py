from __future__ import annotations

from copy import copy
from datetime import datetime
import io
import os
import tempfile

import qrcode
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.drawing.image import Image as XLImage
from openpyxl.utils import get_column_letter
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string


SHEET_NAME = "F-1.03"

CHECK = "✓"
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")
LEFT_WRAP = Alignment(horizontal="left", vertical="center", wrap_text=True)
FONT_CHECK = Font(bold=True, size=11)


# =========================================================
# DEFAULT VALUE
# =========================================================
DEFAULT_DATA = {
    "nomor_kk": "-",
    "nama_lengkap": "-",
    "nik": "-",
    "jenis_pemohon": "skp-wni",  # skp-wni | skpln | sktt-orang-asing

    # asal
    "asal_kode_provinsi": "00",
    "asal_provinsi": "PROVINSI",
    "asal_kode_kabupaten_kota": "0000",
    "asal_kabupaten_kota": "KABUPATEN/KOTA",
    "asal_kode_kecamatan": "000",
    "asal_kecamatan": "KECAMATAN",
    "asal_kode_desa_kelurahan": "0000",
    "asal_desa_kelurahan": "DESA/KELURAHAN",
    "asal_rt": "000",
    "asal_rw": "000",
    "asal_kode_pos": "00000",

    "jenis_perpindahan": "antar-kabupaten-kota",

    # tujuan
    "tujuan_kode_provinsi": "00",
    "tujuan_provinsi": "PROVINSI",
    "tujuan_kode_kabupaten_kota": "0000",
    "tujuan_kabupaten_kota": "KABUPATEN/KOTA",
    "tujuan_kode_kecamatan": "000",
    "tujuan_kecamatan": "KECAMATAN",
    "tujuan_kode_desa_kelurahan": "0000",
    "tujuan_desa_kelurahan": "DESA/KELURAHAN",
    "tujuan_rt": "000",
    "tujuan_rw": "000",
    "tujuan_kode_pos": "00000",

    "alasan_pindah": "pekerjaan",
    "alasan_pindah_lainnya": "",

    "jenis_kepindahan": "kepala-keluarga-sebagian-anggota",
    "status_kk_tidak_pindah": "buat-kk-baru",
    "status_kk_pindah": "buat-kk-baru",

    "daftar_anggota_pindah": [],

    "nama_sponsor": "",
    "tipe_sponsor": "tanpa-sponsor",
    "alamat_sponsor": "",

    "nomor_kitas_kitap": "",
    "masa_berlaku_kitas_kitap": "",

    "negara_tujuan": "",
    "alamat_tujuan_luar_negeri": "",

    "penanggung_jawab": "-",
    "rencana_tanggal_pindah": "",

    "tempat_surat": "SRAGEN",
    "tanggal_surat": "",
    "nama_kades": "KEPALA DESA",
    "nip_kades": "-",
    "qr_ttd": "",
}


SHDK_NUMERIC = {
    "kepala-keluarga": "1",
    "suami": "2",
    "istri": "3",
    "anak": "4",
    "menantu": "5",
    "cucu": "6",
    "orang-tua": "7",
    "mertua": "8",
    "famili-lain": "9",
    "pembantu": "10",
    "lainnya": "11",
}


CHECKBOX_MAP = {
    "jenis_pemohon": {
        "skp-wni": "I13",
        "skpln": "I14",
        "sktt-orang-asing": "I15",
    },
    "jenis_perpindahan": {
        "dalam-desa-kelurahan": "H24",
        "antar-desa-kelurahan": "H25",
        "antar-kecamatan": "H26",
        "antar-kabupaten-kota": "H27",
        "antar-provinsi": "H28",
    },
    "alasan_pindah": {
        "pekerjaan": "H36",
        "pendidikan": "H37",
        "keamanan": "U36",
        "kesehatan": "U37",
        "perumahan": "AE36",
        "keluarga": "AE37",
        "lainnya": "AK36",
    },
    "jenis_kepindahan": {
        "kepala-keluarga": "H39",
        "kepala-keluarga-sebagian-anggota": "U39",
        "kepala-keluarga-seluruh-anggota": "H40",
        "anggota-keluarga": "U40",
    },
    "status_kk_tidak_pindah": {
        "numpang-kk": "H42",
        "buat-kk-baru": "R42",
    },
    "status_kk_pindah": {
        "numpang-kk": "H45",
        "buat-kk-baru": "R45",
    },
    "tipe_sponsor": {
        "organisasi-internasional": "G67",
        "pemerintah": "R67",
        "perusahaan": "Z67",
        "perorangan": "G69",
        "tanpa-sponsor": "R69",
    },
}


TEXT_FIELDS = {
    "nomor_kk": "H6",
    "nama_lengkap": "H8",
    "nik": "H10",

    "asal_rt": "AM18",
    "asal_rw": "AQ18",
    "asal_desa_kelurahan": "T20",
    "asal_kecamatan": "AN20",
    "asal_kabupaten_kota": "T21",
    "asal_provinsi": "AN21",
    "asal_kode_pos": "T22",

    "tujuan_rt": "AM30",
    "tujuan_rw": "AQ30",
    "tujuan_desa_kelurahan": "T32",
    "tujuan_kecamatan": "AN32",
    "tujuan_kabupaten_kota": "T33",
    "tujuan_provinsi": "AN33",
    "tujuan_kode_pos": "T34",

    "nama_sponsor": "G65",
    "alamat_sponsor": "G71",
    "nomor_kitas_kitap": "G73",
    "masa_berlaku_kitas_kitap": "W73",
    "negara_tujuan": "G78",
    "penanggung_jawab": "G84",
}

TABLE_ROW_START = 51
TABLE_ROW_END = 60


# =========================================================
# Helpers
# =========================================================
def normalize_text(value):
    if value is None:
        return ""
    return str(value).strip()


def safe_default(value, default_value=""):
    value = normalize_text(value)
    return value if value else default_value


def sync_wilayah_pair(data: dict, default: dict, kode_key: str, nama_key: str):
    kode = normalize_text(data.get(kode_key))
    nama = normalize_text(data.get(nama_key))

    if not kode and not nama:
        data[kode_key] = default[kode_key]
        data[nama_key] = default[nama_key]
    elif not kode:
        data[kode_key] = default[kode_key]
        data[nama_key] = nama
    elif not nama:
        data[kode_key] = kode
        data[nama_key] = default[nama_key]
    else:
        data[kode_key] = kode
        data[nama_key] = nama


def get_real_cell_for_merged(ws, cell_ref: str) -> str:
    for merged_range in ws.merged_cells.ranges:
        if cell_ref in merged_range:
            return merged_range.start_cell.coordinate
    return cell_ref


def set_value(ws, cell_ref: str, value, align: Alignment | None = None, as_text: bool = True):
    if value is None or value == "":
        return

    real_cell = get_real_cell_for_merged(ws, cell_ref)
    text_value = str(value)
    ws[real_cell] = text_value
    if as_text:
        ws[real_cell].number_format = "@"
    if align:
        ws[real_cell].alignment = copy(align)


def clear_cell(ws, cell_ref: str):
    real_cell = get_real_cell_for_merged(ws, cell_ref)
    ws[real_cell] = None


def clear_cells(ws, cell_refs):
    for ref in cell_refs:
        clear_cell(ws, ref)


def check_cell(ws, cell_ref: str):
    real_cell = get_real_cell_for_merged(ws, cell_ref)
    ws[real_cell] = CHECK
    ws[real_cell].alignment = copy(CENTER)
    ws[real_cell].font = copy(FONT_CHECK)


def create_qr_image_file(text):
    if not text:
        return None

    temp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".png")
    temp_file.close()

    qr = qrcode.QRCode(
        version=2,
        error_correction=qrcode.constants.ERROR_CORRECT_M,
        box_size=8,
        border=1,
    )
    qr.add_data(text)
    qr.make(fit=True)

    img = qr.make_image(fill_color="black", back_color="white")
    img.save(temp_file.name)
    return temp_file.name


def safe_remove_file(path):
    try:
        if path and os.path.exists(path):
            os.remove(path)
    except Exception:
        pass


def insert_qr_anchor(ws, image_path: str, cell: str, width: int = 80, height: int = 80):
    if not image_path or not os.path.exists(image_path):
        return

    img = XLImage(image_path)
    img.width = width
    img.height = height
    img.anchor = cell
    ws.add_image(img)


def mark_one(ws, mapping: dict[str, str], selected_value: str | None):
    clear_cells(ws, mapping.values())
    if selected_value in mapping:
        check_cell(ws, mapping[selected_value])


def parse_date_parts(value) -> tuple[str, str, str]:
    if not value:
        return "", "", ""

    if isinstance(value, datetime):
        dt = value
    else:
        text = str(value).strip()
        dt = None
        for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%Y/%m/%d", "%d/%m/%Y"):
            try:
                dt = datetime.strptime(text, fmt)
                break
            except ValueError:
                pass

        if dt is None:
            digits = "".join(ch for ch in text if ch.isdigit())
            if len(digits) == 8:
                if digits[:4].startswith(("19", "20")):
                    dt = datetime.strptime(digits, "%Y%m%d")
                else:
                    dt = datetime.strptime(digits, "%d%m%Y")
            else:
                return "", "", ""

    return f"{dt.day:02d}", f"{dt.month:02d}", f"{dt.year:04d}"


def fill_digits(ws, cell_refs: list[str], text: str):
    padded = (text or "")[: len(cell_refs)].ljust(len(cell_refs))
    for ref, char in zip(cell_refs, padded):
        set_value(ws, ref, char if char.strip() else "", CENTER)


def split_text_for_lines(text: str, line_lengths: list[int]) -> list[str]:
    if not text:
        return [""] * len(line_lengths)

    text = " ".join(str(text).split())
    words = text.split(" ")
    lines: list[str] = []
    current = ""

    for max_len in line_lengths:
        if not words and not current:
            lines.append("")
            continue

        while words:
            candidate = words[0] if not current else f"{current} {words[0]}"
            if len(candidate) <= max_len:
                current = candidate
                words.pop(0)
            else:
                break

        if not current and words:
            current = words.pop(0)[:max_len]

        lines.append(current)
        current = ""

    if words and lines:
        remainder = " ".join(words)
        lines[-1] = (lines[-1] + " " + remainder).strip()[: line_lengths[-1]]

    return lines


# =========================================================
# Transform raw data
# =========================================================
def transform_raw_to_data(raw_data: dict | None) -> dict:
    raw_data = raw_data or {}

    data = dict(DEFAULT_DATA)
    data.update(raw_data)

    # utama
    data["nomor_kk"] = safe_default(data.get("nomor_kk"), DEFAULT_DATA["nomor_kk"])
    data["nama_lengkap"] = safe_default(data.get("nama_lengkap"), DEFAULT_DATA["nama_lengkap"])
    data["nik"] = safe_default(data.get("nik"), DEFAULT_DATA["nik"])
    data["jenis_pemohon"] = safe_default(data.get("jenis_pemohon"), DEFAULT_DATA["jenis_pemohon"])

    # sinkron wilayah asal
    sync_wilayah_pair(data, DEFAULT_DATA, "asal_kode_provinsi", "asal_provinsi")
    sync_wilayah_pair(data, DEFAULT_DATA, "asal_kode_kabupaten_kota", "asal_kabupaten_kota")
    sync_wilayah_pair(data, DEFAULT_DATA, "asal_kode_kecamatan", "asal_kecamatan")
    sync_wilayah_pair(data, DEFAULT_DATA, "asal_kode_desa_kelurahan", "asal_desa_kelurahan")

    # sinkron wilayah tujuan
    sync_wilayah_pair(data, DEFAULT_DATA, "tujuan_kode_provinsi", "tujuan_provinsi")
    sync_wilayah_pair(data, DEFAULT_DATA, "tujuan_kode_kabupaten_kota", "tujuan_kabupaten_kota")
    sync_wilayah_pair(data, DEFAULT_DATA, "tujuan_kode_kecamatan", "tujuan_kecamatan")
    sync_wilayah_pair(data, DEFAULT_DATA, "tujuan_kode_desa_kelurahan", "tujuan_desa_kelurahan")

    # field alamat asal/tujuan yang memang ada di spreadsheet
    for key in [
        "asal_rt", "asal_rw", "asal_kode_pos",
        "tujuan_rt", "tujuan_rw", "tujuan_kode_pos",
        "jenis_perpindahan", "alasan_pindah", "alasan_pindah_lainnya",
        "jenis_kepindahan", "status_kk_tidak_pindah", "status_kk_pindah",
        "nama_sponsor", "tipe_sponsor", "alamat_sponsor",
        "nomor_kitas_kitap", "masa_berlaku_kitas_kitap",
        "negara_tujuan", "alamat_tujuan_luar_negeri",
        "penanggung_jawab", "rencana_tanggal_pindah",
        "tempat_surat", "tanggal_surat",
    ]:
        if key in DEFAULT_DATA:
            data[key] = safe_default(data.get(key), DEFAULT_DATA[key])
        else:
            data[key] = safe_default(data.get(key))

    # kades
    nama_kades_raw = raw_data.get("nama_kades")
    if not normalize_text(nama_kades_raw):
        nama_kades_raw = raw_data.get("nama_kepala_desa")

    nip_kades_raw = raw_data.get("nip_kades")
    if not normalize_text(nip_kades_raw):
        nip_kades_raw = raw_data.get("nip_kepala_desa")

    data["nama_kades"] = safe_default(nama_kades_raw, DEFAULT_DATA["nama_kades"])
    data["nip_kades"] = safe_default(nip_kades_raw, DEFAULT_DATA["nip_kades"])

    qr_text = safe_default(raw_data.get("qr_ttd"))
    if not qr_text:
        qr_text = f"Ditandatangani secara elektronik oleh {data['nama_kades']}"
        if data["nip_kades"] and data["nip_kades"] != "-":
            qr_text += f", NIP {data['nip_kades']}"
    data["qr_text"] = qr_text

    # anggota pindah
    anggota_raw = raw_data.get("daftar_anggota_pindah")
    if not isinstance(anggota_raw, list):
        anggota_raw = []

    anggota_normalized = []
    for item in anggota_raw:
        if not isinstance(item, dict):
            continue

        anggota_normalized.append({
            "anggota_nik": item.get("anggota_nik") or item.get("nik", ""),
            "anggota_nama_lengkap": item.get("anggota_nama_lengkap") or item.get("nama", ""),
            "anggota_shdk": item.get("anggota_shdk") or item.get("status", ""),
        })

    data["daftar_anggota_pindah"] = anggota_normalized

    return data


# =========================================================
# Fillers
# =========================================================
def fill_main_text_fields(ws, data: dict):
    for key, cell_ref in TEXT_FIELDS.items():
        set_value(ws, cell_ref, data.get(key, ""), LEFT)


def fill_checkboxes(ws, data: dict):
    for field_name, mapping in CHECKBOX_MAP.items():
        mark_one(ws, mapping, data.get(field_name))


def fill_alasan_lainnya(ws, data: dict):
    value = data.get("alasan_pindah_lainnya", "") if data.get("alasan_pindah") == "lainnya" else ""
    set_value(ws, "AL37", value, LEFT)


def fill_alamat_tujuan_luar_negeri(ws, data: dict):
    line1, line2 = split_text_for_lines(data.get("alamat_tujuan_luar_negeri", ""), [110, 110])
    set_value(ws, "G80", line1, LEFT)
    set_value(ws, "G82", line2, LEFT)


def fill_anggota_table(ws, data: dict):
    anggota_list = data.get("daftar_anggota_pindah", [])[: TABLE_ROW_END - TABLE_ROW_START + 1]

    for row in range(TABLE_ROW_START, TABLE_ROW_END + 1):
        clear_cells(ws, [f"C{row}", f"F{row}", f"X{row}", f"AX{row}"])

    for idx, anggota in enumerate(anggota_list, start=1):
        row = TABLE_ROW_START + idx - 1

        set_value(ws, f"C{row}", idx, CENTER)
        set_value(ws, f"F{row}", anggota.get("anggota_nik", ""), LEFT)
        set_value(ws, f"X{row}", anggota.get("anggota_nama_lengkap", ""), LEFT)

        shdk_raw = anggota.get("anggota_shdk", "")
        shdk = SHDK_NUMERIC.get(shdk_raw, shdk_raw)

        set_value(ws, f"AX{row}", shdk, CENTER)


def fill_rencana_pindah(ws, data: dict):
    day, month, year = parse_date_parts(data.get("rencana_tanggal_pindah"))
    fill_digits(ws, ["J86", "K86"], day)
    fill_digits(ws, ["Q86", "R86"], month)
    fill_digits(ws, ["X86", "Y86", "Z86", "AA86"], year)


def fill_footer(ws, data: dict):
    tempat = data.get("tempat_surat", "") or data.get("tempat_ttd", "")
    tgl = data.get("tanggal_surat", "") or data.get("tanggal_ttd", "")
    day, month, year = parse_date_parts(tgl)

    if tempat:
        set_value(ws, "AK91", tempat, LEFT)

    month_names = {
        "01": "Januari", "02": "Februari", "03": "Maret", "04": "April",
        "05": "Mei", "06": "Juni", "07": "Juli", "08": "Agustus",
        "09": "September", "10": "Oktober", "11": "November", "12": "Desember",
    }

    if day or month or year:
        bulan_text = month_names.get(month, month)
        tanggal_full = f"{day} {bulan_text} {year}".strip()
        set_value(ws, "AN91", tanggal_full, LEFT)

    if data.get("nama_kades"):
        set_value(ws, "Y99", data["nama_kades"], CENTER)

    if data.get("nip_kades"):
        set_value(ws, "Y101", data["nip_kades"], CENTER)

    if data.get("nama_lengkap"):
        set_value(ws, "AL99", data["nama_lengkap"], CENTER)

    qr_path = create_qr_image_file(data.get("qr_text", ""))
    if qr_path:
        insert_qr_anchor(ws, qr_path, "AA94", width=120, height=120)

    return qr_path


def fill_document(ws, data: dict):
    fill_main_text_fields(ws, data)
    fill_checkboxes(ws, data)
    fill_alasan_lainnya(ws, data)
    fill_alamat_tujuan_luar_negeri(ws, data)
    fill_anggota_table(ws, data)
    fill_rencana_pindah(ws, data)
    qr_path = fill_footer(ws, data)
    return qr_path


# =========================================================
# Main generator
# =========================================================
def generate_b09_excel(template_path: str, raw_data: dict | None):
    data = transform_raw_to_data(raw_data)

    wb = load_workbook(template_path)
    ws = wb[SHEET_NAME]

    qr_path = fill_document(ws, data)

    output = io.BytesIO()
    try:
        wb.save(output)
        output.seek(0)
        return output.getvalue()
    finally:
        safe_remove_file(qr_path)