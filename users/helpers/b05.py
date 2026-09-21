from copy import copy
from datetime import datetime
import io
import os
import tempfile

import qrcode
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
from openpyxl.utils import get_column_letter


SHEET_NAME = "Permohonan KTP"

CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
CHECK = "✓"
FONT_CHECK = Font(bold=True)


# =========================================================
# DEFAULT VALUE
# =========================================================
DEFAULT_DATA = {
    # wilayah
    "kode_provinsi": "33",
    "nama_provinsi": "JAWA TENGAH",

    "kode_kabupaten_kota": "3374",
    "nama_kabupaten_kota": "KOTA SEMARANG",

    "kode_kecamatan": "000",
    "nama_kecamatan": "KECAMATAN",

    "kode_kelurahan_desa": "0000",
    "nama_kelurahan_desa": "KELURAHAN/DESA",

    # jenis permohonan
    "jenis_permohonan_ktp": "baru",  # baru | perpanjangan | penggantian

    # pemohon
    "nama_lengkap": "-",
    "nik": "-",
    "nomor_kk": "-",
    "nomor_telepon": "-",
    "alamat": "-",
    "rt": "000",
    "rw": "000",
    "kode_pos": "00000",

    # footer / tanda tangan
    "tempat": "SRAGEN",
    "tanggal": "",
    "bulan": "",
    "tahun": "",
    "tanggal_ttd": "",
    "tempat_ttd": "SRAGEN",

    "nama_kades": "KEPALA DESA",
    "nip_kades": "-",
    "qr_ttd": "",
}


# =========================================================
# HELPER DASAR
# =========================================================
def normalize_text(value):
    if value is None:
        return ""
    return str(value).strip()


def safe_default(value, default_value=""):
    value = normalize_text(value)
    return value if value else default_value


def sync_wilayah(data, default):
    pairs = [
        ("kode_provinsi", "nama_provinsi"),
        ("kode_kabupaten_kota", "nama_kabupaten_kota"),
        ("kode_kecamatan", "nama_kecamatan"),
        ("kode_kelurahan_desa", "nama_kelurahan_desa"),
    ]

    for kode_key, nama_key in pairs:
        kode = normalize_text(data.get(kode_key))
        nama = normalize_text(data.get(nama_key))

        if not kode and not nama:
            data[kode_key] = default[kode_key]
            data[nama_key] = default[nama_key]
        elif not kode:
            data[kode_key] = default[kode_key]
            data[nama_key] = nama
        elif not nama:
            data[kode_key] = kode
            data[nama_key] = default[nama_key]
        else:
            data[kode_key] = kode
            data[nama_key] = nama


def build_footer_date_parts(raw_data):
    tempat = safe_default(raw_data.get("tempat"), DEFAULT_DATA["tempat"])
    tanggal = safe_default(raw_data.get("tanggal"))
    bulan = safe_default(raw_data.get("bulan"))
    tahun = safe_default(raw_data.get("tahun"))

    if not (tanggal and bulan and tahun):
        tanggal_ttd = normalize_text(raw_data.get("tanggal_ttd"))
        if tanggal_ttd:
            for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%Y/%m/%d", "%d/%m/%Y"):
                try:
                    parsed = datetime.strptime(tanggal_ttd, fmt)
                    tanggal = str(parsed.day)
                    bulan = parsed.strftime("%B")
                    tahun = str(parsed.year)
                    break
                except ValueError:
                    continue

    return tempat, tanggal, bulan, tahun


def build_footer_date_text(data):
    parts = [data.get("tempat", ""), data.get("tanggal", ""), data.get("bulan", ""), data.get("tahun", "")]
    tempat = normalize_text(parts[0])
    tanggal = normalize_text(parts[1])
    bulan = normalize_text(parts[2])
    tahun = normalize_text(parts[3])

    if tempat or tanggal or bulan or tahun:
        return f"{tempat}, {tanggal} {bulan} {tahun}".strip().strip(",")
    return ""


def get_real_cell_for_merged(ws, cell_ref):
    for merged_range in ws.merged_cells.ranges:
        if cell_ref in merged_range:
            return merged_range.start_cell.coordinate
    return cell_ref


def set_value(ws, cell_ref, value, align=None):
    if value is None or value == "":
        return

    real_cell = get_real_cell_for_merged(ws, cell_ref)
    ws[real_cell] = value

    if align:
        ws[real_cell].alignment = copy(align)


def check_cell(ws, cell_ref):
    real_cell = get_real_cell_for_merged(ws, cell_ref)
    ws[real_cell] = CHECK
    ws[real_cell].alignment = copy(CENTER)
    ws[real_cell].font = copy(FONT_CHECK)


def create_qr_image_file(text):
    if not text:
        return None

    temp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".png")
    temp_file.close()

    qr = qrcode.QRCode(
        version=2,
        error_correction=qrcode.constants.ERROR_CORRECT_M,
        box_size=8,
        border=1,
    )
    qr.add_data(text)
    qr.make(fit=True)

    img = qr.make_image(fill_color="black", back_color="white")
    img.save(temp_file.name)
    return temp_file.name


def safe_remove_file(path):
    try:
        if path and os.path.exists(path):
            os.remove(path)
    except Exception:
        pass


def insert_qr(ws, image_path, cell="AE40"):
    if not image_path or not os.path.exists(image_path):
        return

    img = XLImage(image_path)
    img.width = 70
    img.height = 70
    img.anchor = cell

    ws.add_image(img)


# =========================================================
# TRANSFORM RAW DATA
# =========================================================
def transform_raw_to_data(raw_data):
    raw_data = raw_data or {}

    data = dict(DEFAULT_DATA)
    data.update(raw_data)

    # wilayah
    sync_wilayah(data, DEFAULT_DATA)

    # pemohon
    data["jenis_permohonan_ktp"] = safe_default(data.get("jenis_permohonan_ktp"), DEFAULT_DATA["jenis_permohonan_ktp"])
    data["nama_lengkap"] = safe_default(data.get("nama_lengkap"), DEFAULT_DATA["nama_lengkap"])
    data["nik"] = safe_default(data.get("nik"), DEFAULT_DATA["nik"])
    data["nomor_kk"] = safe_default(data.get("nomor_kk"), DEFAULT_DATA["nomor_kk"])
    data["nomor_telepon"] = safe_default(data.get("nomor_telepon"), DEFAULT_DATA["nomor_telepon"])
    data["alamat"] = safe_default(data.get("alamat"), DEFAULT_DATA["alamat"])
    data["rt"] = safe_default(data.get("rt"), DEFAULT_DATA["rt"])
    data["rw"] = safe_default(data.get("rw"), DEFAULT_DATA["rw"])
    data["kode_pos"] = safe_default(data.get("kode_pos"), DEFAULT_DATA["kode_pos"])

    # kades
    nama_kades_raw = raw_data.get("nama_kades")
    if not normalize_text(nama_kades_raw):
        nama_kades_raw = raw_data.get("nama_kepala_desa")

    nip_kades_raw = raw_data.get("nip_kades")
    if not normalize_text(nip_kades_raw):
        nip_kades_raw = raw_data.get("nip_kepala_desa")

    data["nama_kades"] = safe_default(nama_kades_raw, DEFAULT_DATA["nama_kades"])
    data["nip_kades"] = safe_default(nip_kades_raw, DEFAULT_DATA["nip_kades"])

    # footer tanggal
    tempat, tanggal, bulan, tahun = build_footer_date_parts(raw_data)
    data["tempat"] = tempat
    data["tanggal"] = tanggal
    data["bulan"] = bulan
    data["tahun"] = tahun

    # QR
    qr_text = safe_default(raw_data.get("qr_ttd"))
    if not qr_text:
        qr_text = f"Ditandatangani secara elektronik oleh {data['nama_kades']}"
        if data["nip_kades"] and data["nip_kades"] != "-":
            qr_text += f", NIP {data['nip_kades']}"
    data["qr_text"] = qr_text

    return data


# =========================================================
# FILL HEADER
# =========================================================
def fill_header_wilayah(ws, data):
    # kode wilayah
    set_value(ws, "M9", data["kode_provinsi"], LEFT)
    set_value(ws, "M10", data["kode_kabupaten_kota"], LEFT)
    set_value(ws, "M11", data["kode_kecamatan"], LEFT)
    set_value(ws, "M12", data["kode_kelurahan_desa"], LEFT)

    # nama wilayah
    set_value(ws, "R9", data["nama_provinsi"], LEFT)
    set_value(ws, "R10", data["nama_kabupaten_kota"], LEFT)
    set_value(ws, "R11", data["nama_kecamatan"], LEFT)
    set_value(ws, "R12", data["nama_kelurahan_desa"], LEFT)


# =========================================================
# JENIS PERMOHONAN
# =========================================================
def fill_jenis_permohonan(ws, data):
    jenis = normalize_text(data.get("jenis_permohonan_ktp", "")).lower()

    if jenis == "baru":
        check_cell(ws, "H14")
    elif jenis == "perpanjangan":
        check_cell(ws, "N14")
    elif jenis == "penggantian":
        check_cell(ws, "V14")


# =========================================================
# DATA PEMOHON
# =========================================================
def fill_data_pemohon(ws, data):
    set_value(ws, "F17", data["nama_lengkap"], LEFT)
    set_value(ws, "F19", data["nomor_kk"], LEFT)   
    set_value(ws, "F21", data["nik"], LEFT)        

    set_value(ws, "F23", data["alamat"], LEFT)
    set_value(ws, "I27", data["rt"], CENTER)
    set_value(ws, "P27", data["rw"], CENTER)
    set_value(ws, "AA27", data["kode_pos"], CENTER)


# =========================================================
# FOOTER
# =========================================================
def fill_footer(ws, data):
    tanggal_lokasi = build_footer_date_text(data)
    set_value(ws, "X29", tanggal_lokasi, CENTER)

    # nama pemohon
    set_value(ws, "AE35", data["nama_lengkap"], CENTER)
    # Kepala Desa / Lurah
    set_value(ws, "W44", data["nama_kades"], CENTER)
    set_value(ws, "W45", f'NIP. {data["nip_kades"]}', CENTER)

    qr_path = create_qr_image_file(data.get("qr_text", ""))
    if qr_path:
        insert_qr(ws, qr_path, "AD39")

    return qr_path


# =========================================================
# MAIN GENERATOR
# =========================================================
def generate_b05_excel(template_path, raw_data):
    raw_data = raw_data or {}

    data = transform_raw_to_data(raw_data)

    wb = load_workbook(template_path)
    ws = wb[SHEET_NAME]

    fill_header_wilayah(ws, data)
    fill_jenis_permohonan(ws, data)
    fill_data_pemohon(ws, data)

    qr_path = fill_footer(ws, data)

    output = io.BytesIO()
    try:
        wb.save(output)
        output.seek(0)
        return output.getvalue()
    finally:
        safe_remove_file(qr_path)