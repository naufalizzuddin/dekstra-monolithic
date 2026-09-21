from copy import copy
from datetime import datetime
import io
import os
import tempfile

import qrcode
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
from openpyxl.utils import get_column_letter


SHEET_NAME = "Formulir KK"

CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
CHECK = "✓"
FONT_CHECK = Font(bold=True)


# =========================================================
# DEFAULT VALUE
# =========================================================
DEFAULT_DATA = {
    # wilayah
    "kode_provinsi": "33",
    "nama_provinsi": "JAWA TENGAH",

    "kode_kabupaten_kota": "3374",
    "nama_kabupaten_kota": "KOTA SEMARANG",

    "kode_kecamatan": "000",
    "nama_kecamatan": "KECAMATAN",

    "kode_kelurahan_desa": "0000",
    "nama_kelurahan_desa": "KELURAHAN/DESA",

    # data pemohon / KK baru
    "nama_lengkap": "-",
    "nik": "-",
    "nama_kepala_keluarga": "-",
    "nomor_kk": "-",
    "alamat": "-",
    "rt": "000",
    "rw": "000",
    "desa_kelurahan": "-",
    "kecamatan": "-",
    "kabupaten_kota": "-",
    "provinsi": "-",
    "kode_pos": "00000",
    "nomor_telepon": "-",

    # data keluarga lama
    "nama_kepala_keluarga_lama": "-",
    "nomor_kk_lama": "-",
    "alamat_lama": "-",
    "rt_lama": "000",
    "rw_lama": "000",
    "desa_kelurahan_lama": "-",
    "kecamatan_lama": "-",
    "kabupaten_kota_lama": "-",
    "provinsi_lama": "-",
    "kode_pos_lama": "00000",
    "nomor_telepon_lama": "-",

    # alasan dan jumlah
    "alasan_permohonan": "",
    "jumlah_anggota_keluarga": "0",

    # anggota keluarga
    "anggota_keluraga": [],
    "anggota_keluarga": [],

    # footer / ttd
    "tempat": "SRAGEN",
    "tanggal": "",
    "bulan": "",
    "tahun": "",
    "tanggal_ttd": "",
    "tempat_ttd": "SRAGEN",

    "nama_kades": "KEPALA DESA",
    "nip_kades": "-",
    "qr_ttd": "",

    "tanggal_pemasukan_tgl": "",
    "tanggal_pemasukan_bln": "",
    "tanggal_pemasukan_thn": "",
}


SHDK_MAP = {
    "kepala-keluarga": "01",
    "suami": "02",
    "istri": "03",
    "anak": "04",
    "menantu": "05",
    "cucu": "06",
    "orang-tua": "07",
    "mertua": "08",
    "famili-lain": "09",
    "pembantu": "10",
    "lainnya": "11",
}


# =========================================================
# HELPER DASAR
# =========================================================
def normalize_text(value):
    if value is None:
        return ""
    return str(value).strip()


def safe_default(value, default_value=""):
    value = normalize_text(value)
    return value if value else default_value


def normalize_date(value):
    value = normalize_text(value)
    if not value:
        return ""

    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%Y/%m/%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(value, fmt).strftime("%d-%m-%Y")
        except ValueError:
            continue

    return value


def build_footer_date_parts(raw_data):
    tempat = safe_default(raw_data.get("tempat"), DEFAULT_DATA["tempat"])
    tanggal = safe_default(raw_data.get("tanggal"))
    bulan = safe_default(raw_data.get("bulan"))
    tahun = safe_default(raw_data.get("tahun"))

    if not (tanggal and bulan and tahun):
        tanggal_ttd = normalize_text(raw_data.get("tanggal_ttd"))
        if tanggal_ttd:
            for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%Y/%m/%d", "%d/%m/%Y"):
                try:
                    parsed = datetime.strptime(tanggal_ttd, fmt)
                    tanggal = str(parsed.day)
                    bulan = parsed.strftime("%B")
                    tahun = str(parsed.year)
                    break
                except ValueError:
                    continue

    return tempat, tanggal, bulan, tahun


def sync_wilayah(data, default):
    pairs = [
        ("kode_provinsi", "nama_provinsi"),
        ("kode_kabupaten_kota", "nama_kabupaten_kota"),
        ("kode_kecamatan", "nama_kecamatan"),
        ("kode_kelurahan_desa", "nama_kelurahan_desa"),
    ]

    for kode_key, nama_key in pairs:
        kode = normalize_text(data.get(kode_key))
        nama = normalize_text(data.get(nama_key))

        if not kode and not nama:
            data[kode_key] = default[kode_key]
            data[nama_key] = default[nama_key]
        elif not kode:
            data[kode_key] = default[kode_key]
            data[nama_key] = nama
        elif not nama:
            data[kode_key] = kode
            data[nama_key] = default[nama_key]
        else:
            data[kode_key] = kode
            data[nama_key] = nama


def get_real_cell_for_merged(ws, cell_ref):
    for merged_range in ws.merged_cells.ranges:
        if cell_ref in merged_range:
            return merged_range.start_cell.coordinate
    return cell_ref


def set_value(ws, cell_ref, value, align=None):
    if value is None or value == "":
        return
    real_cell = get_real_cell_for_merged(ws, cell_ref)
    ws[real_cell] = value
    if align:
        ws[real_cell].alignment = copy(align)


def check_cell(ws, cell_ref):
    real_cell = get_real_cell_for_merged(ws, cell_ref)
    ws[real_cell] = CHECK
    ws[real_cell].alignment = copy(CENTER)
    ws[real_cell].font = copy(FONT_CHECK)


def get_shdk_code(status_value):
    if not status_value:
        return ""
    return SHDK_MAP.get(normalize_text(status_value).lower(), "11")


def create_qr_image_file(text):
    if not text:
        return None

    temp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".png")
    temp_file.close()

    qr = qrcode.QRCode(
        version=2,
        error_correction=qrcode.constants.ERROR_CORRECT_M,
        box_size=8,
        border=1,
    )
    qr.add_data(text)
    qr.make(fit=True)

    img = qr.make_image(fill_color="black", back_color="white")
    img.save(temp_file.name)
    return temp_file.name


def safe_remove_file(path):
    try:
        if path and os.path.exists(path):
            os.remove(path)
    except Exception:
        pass


def insert_qr_center_range(ws, image_path, start_cell, end_cell, width=65, height=65):
    if not image_path or not os.path.exists(image_path):
        return

    img = XLImage(image_path)
    img.width = width
    img.height = height

    start_col_letter, start_row = coordinate_from_string(start_cell)
    end_col_letter, end_row = coordinate_from_string(end_cell)

    start_col = column_index_from_string(start_col_letter)
    end_col = column_index_from_string(end_col_letter)

    total_width = 0
    for col in range(start_col, end_col + 1):
        col_letter = get_column_letter(col)
        total_width += int((ws.column_dimensions[col_letter].width or 8.43) * 7)

    total_height = 0
    for row in range(start_row, end_row + 1):
        total_height += int((ws.row_dimensions[row].height or 15) * 1.33)

    offset_x = max((total_width - width) // 2, 0)
    offset_y = max((total_height - height) // 2, 0)

    marker = AnchorMarker(
        col=start_col - 1,
        colOff=offset_x * 9525,
        row=start_row - 1,
        rowOff=offset_y * 9525,
    )

    size = XDRPositiveSize2D(
        cx=width * 9525,
        cy=height * 9525,
    )

    img.anchor = OneCellAnchor(_from=marker, ext=size)
    ws.add_image(img)


# =========================================================
# TRANSFORM RAW DATA
# =========================================================
def transform_raw_to_data(raw_data):
    raw_data = raw_data or {}

    data = dict(DEFAULT_DATA)
    data.update(raw_data)

    # wilayah
    sync_wilayah(data, DEFAULT_DATA)

    # data pemohon / KK baru
    data["nama_lengkap"] = safe_default(data.get("nama_lengkap"), DEFAULT_DATA["nama_lengkap"])
    data["nik"] = safe_default(data.get("nik"), DEFAULT_DATA["nik"])
    data["nama_kepala_keluarga"] = safe_default(data.get("nama_kepala_keluarga"), DEFAULT_DATA["nama_kepala_keluarga"])
    data["nomor_kk"] = safe_default(data.get("nomor_kk"), DEFAULT_DATA["nomor_kk"])
    data["alamat"] = safe_default(data.get("alamat"), DEFAULT_DATA["alamat"])
    data["rt"] = safe_default(data.get("rt"), DEFAULT_DATA["rt"])
    data["rw"] = safe_default(data.get("rw"), DEFAULT_DATA["rw"])
    data["desa_kelurahan"] = safe_default(data.get("desa_kelurahan"), DEFAULT_DATA["desa_kelurahan"])
    data["kecamatan"] = safe_default(data.get("kecamatan"), DEFAULT_DATA["kecamatan"])
    data["kabupaten_kota"] = safe_default(data.get("kabupaten_kota"), DEFAULT_DATA["kabupaten_kota"])
    data["provinsi"] = safe_default(data.get("provinsi"), DEFAULT_DATA["provinsi"])
    data["kode_pos"] = safe_default(data.get("kode_pos"), DEFAULT_DATA["kode_pos"])
    data["nomor_telepon"] = safe_default(data.get("nomor_telepon"), DEFAULT_DATA["nomor_telepon"])

    # data keluarga lama
    data["nama_kepala_keluarga_lama"] = safe_default(data.get("nama_kepala_keluarga_lama"), DEFAULT_DATA["nama_kepala_keluarga_lama"])
    data["nomor_kk_lama"] = safe_default(data.get("nomor_kk_lama"), DEFAULT_DATA["nomor_kk_lama"])
    data["alamat_lama"] = safe_default(data.get("alamat_lama"), DEFAULT_DATA["alamat_lama"])
    data["rt_lama"] = safe_default(data.get("rt_lama"), DEFAULT_DATA["rt_lama"])
    data["rw_lama"] = safe_default(data.get("rw_lama"), DEFAULT_DATA["rw_lama"])
    data["desa_kelurahan_lama"] = safe_default(data.get("desa_kelurahan_lama"), DEFAULT_DATA["desa_kelurahan_lama"])
    data["kecamatan_lama"] = safe_default(data.get("kecamatan_lama"), DEFAULT_DATA["kecamatan_lama"])
    data["kabupaten_kota_lama"] = safe_default(data.get("kabupaten_kota_lama"), DEFAULT_DATA["kabupaten_kota_lama"])
    data["provinsi_lama"] = safe_default(data.get("provinsi_lama"), DEFAULT_DATA["provinsi_lama"])
    data["kode_pos_lama"] = safe_default(data.get("kode_pos_lama"), DEFAULT_DATA["kode_pos_lama"])
    data["nomor_telepon_lama"] = safe_default(data.get("nomor_telepon_lama"), DEFAULT_DATA["nomor_telepon_lama"])

    data["alasan_permohonan"] = safe_default(data.get("alasan_permohonan"))
    data["jumlah_anggota_keluarga"] = data.get("jumlah_anggota_keluarga") or DEFAULT_DATA["jumlah_anggota_keluarga"]

    # kades
    nama_kades_raw = raw_data.get("nama_kades")
    if not normalize_text(nama_kades_raw):
        nama_kades_raw = raw_data.get("nama_kepala_desa")

    nip_kades_raw = raw_data.get("nip_kades")
    if not normalize_text(nip_kades_raw):
        nip_kades_raw = raw_data.get("nip_kepala_desa")

    data["nama_kades"] = safe_default(nama_kades_raw, DEFAULT_DATA["nama_kades"])
    data["nip_kades"] = safe_default(nip_kades_raw, DEFAULT_DATA["nip_kades"])

    # footer tanggal
    tempat, tanggal, bulan, tahun = build_footer_date_parts(raw_data)
    data["tempat"] = tempat
    data["tanggal"] = tanggal
    data["bulan"] = bulan
    data["tahun"] = tahun

    # QR
    qr_text = safe_default(raw_data.get("qr_ttd"))
    if not qr_text:
        qr_text = f"Ditandatangani secara elektronik oleh {data['nama_kades']}"
        if data["nip_kades"] and data["nip_kades"] != "-":
            qr_text += f", NIP {data['nip_kades']}"
    data["qr_text"] = qr_text

    # anggota keluarga
    anggota_raw = raw_data.get("anggota_keluarga", None)
    if not isinstance(anggota_raw, list):
        anggota_raw = raw_data.get("anggota_keluraga", None)
    if not isinstance(anggota_raw, list):
        anggota_raw = []

    anggota_normalized = []
    for item in anggota_raw:
        if not isinstance(item, dict):
            continue

        anggota_normalized.append({
            "anggota_nik": item.get("anggota_nik") or item.get("nik", ""),
            "anggota_nama_lengkap": item.get("anggota_nama_lengkap") or item.get("nama", ""),
            "anggota_status_hubungan_keluarga": (
                item.get("anggota_status_hubungan_keluarga") or item.get("status", "")
            ),
        })

    data["anggota_keluraga"] = anggota_normalized

    return data


def fill_header(ws, data):
    # PROVINSI
    set_value(ws, "N10", data["kode_provinsi"], LEFT)
    set_value(ws, "S10", data["nama_provinsi"], LEFT)

    # KABUPATEN / KOTA
    set_value(ws, "N11", data["kode_kabupaten_kota"], LEFT)
    set_value(ws, "S11", data["nama_kabupaten_kota"], LEFT)

    # KECAMATAN
    set_value(ws, "N12", data["kode_kecamatan"], LEFT)
    set_value(ws, "S12", data["nama_kecamatan"], LEFT)

    # KELURAHAN / DESA
    set_value(ws, "N13", data["kode_kelurahan_desa"], LEFT)
    set_value(ws, "S13", data["nama_kelurahan_desa"], LEFT)


def fill_pemohon_baru(ws, data):
    set_value(ws, "N17", data["nama_lengkap"], LEFT)
    set_value(ws, "N19", data["nik"], LEFT)
    set_value(ws, "N21", data["nama_kepala_keluarga"], LEFT)
    set_value(ws, "N23", data["nomor_kk"], LEFT)

    set_value(ws, "N25", data["alamat"], LEFT)
    set_value(ws, "AJ25", data["rt"], CENTER)
    set_value(ws, "AR25", data["rw"], CENTER)

    set_value(ws, "P27", data["desa_kelurahan"], LEFT)
    set_value(ws, "AI27", data["kecamatan"], LEFT)
    set_value(ws, "P29", data["kabupaten_kota"], LEFT)
    set_value(ws, "AI29", data["provinsi"], LEFT)
    set_value(ws, "P31", data["kode_pos"], CENTER)
    set_value(ws, "AI31", data["nomor_telepon"], LEFT)


def fill_keluarga_lama(ws, data):
    set_value(ws, "N34", data["nama_kepala_keluarga_lama"], LEFT)
    set_value(ws, "N36", data["nomor_kk_lama"], LEFT)

    set_value(ws, "N38", data["alamat_lama"], LEFT)
    set_value(ws, "AJ38", data["rt_lama"], CENTER)
    set_value(ws, "AR38", data["rw_lama"], CENTER)

    set_value(ws, "P40", data["desa_kelurahan_lama"], LEFT)
    set_value(ws, "AI40", data["kecamatan_lama"], LEFT)
    set_value(ws, "P42", data["kabupaten_kota_lama"], LEFT)
    set_value(ws, "AI42", data["provinsi_lama"], LEFT)
    set_value(ws, "P44", data["kode_pos_lama"], CENTER)
    set_value(ws, "AI44", data.get("nomor_telepon_lama", ""), LEFT)


def fill_alasan(ws, data):
    alasan = data["alasan_permohonan"]
    if alasan == "penambahan-anggota-keluarga":
        check_cell(ws, "J46")
    elif alasan == "pengurangan-anggota-keluarga":
        check_cell(ws, "J48")
    else:
        check_cell(ws, "AC46")


def fill_jumlah_anggota(ws, data):
    set_value(ws, "M50", data["jumlah_anggota_keluarga"], CENTER)


def fill_anggota(ws, data):
    """
    Mapping tabel anggota:
    row awal 57, lalu gap 1 row -> 57, 59, 61, ...
    Kolom:
    No   = C
    NIK  = F
    Nama = W
    SHDK = AV
    """
    start_row = 57
    row_gap = 2

    anggota_list = data.get("anggota_keluraga", [])

    for i, anggota in enumerate(anggota_list):
        row = start_row + (i * row_gap)

        set_value(ws, f"C{row}", i + 1, CENTER)
        set_value(ws, f"F{row}", anggota.get("anggota_nik", ""), LEFT)
        set_value(ws, f"W{row}", anggota.get("anggota_nama_lengkap", ""), LEFT)
        set_value(
            ws,
            f"AV{row}",
            get_shdk_code(anggota.get("anggota_status_hubungan_keluarga", "")),
            CENTER,
        )


def fill_footer(ws, data):
    tanggal_lokasi = f'{data["tempat"]}, {data["tanggal"]} {data["bulan"]} {data["tahun"]}'.strip().strip(",")
    set_value(ws, "AJ74", tanggal_lokasi, CENTER)

    set_value(ws, "AB84", data["nama_kades"], CENTER)
    set_value(ws, "AL81", data["nama_lengkap"], CENTER)
    set_value(ws, "AB85", f'NIP {data["nip_kades"]}', CENTER)

    qr_path = create_qr_image_file(data.get("qr_text", ""))
    if qr_path:
        insert_qr_center_range(ws, qr_path, "AC79", "AE82", width=65, height=65)

    return qr_path


def generate_b04_excel(template_path, raw_data):
    raw_data = raw_data or {}

    data = transform_raw_to_data(raw_data)

    wb = load_workbook(template_path)
    ws = wb[SHEET_NAME]

    fill_header(ws, data)
    fill_pemohon_baru(ws, data)
    fill_keluarga_lama(ws, data)
    fill_alasan(ws, data)
    fill_jumlah_anggota(ws, data)
    fill_anggota(ws, data)

    qr_path = fill_footer(ws, data)

    output = io.BytesIO()
    try:
        wb.save(output)
        output.seek(0)
        return output.getvalue()
    finally:
        safe_remove_file(qr_path)