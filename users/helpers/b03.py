from copy import copy
from datetime import datetime
import io
import os
import tempfile

import qrcode
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import AnchorMarker, OneCellAnchor
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.utils.cell import coordinate_from_string, column_index_from_string
from openpyxl.utils import get_column_letter


SHEET_NAME = "Formulir KK"

CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)
CHECK = "✓"
FONT_CHECK = Font(bold=True)


# =========================================================
# DEFAULT VALUE
# =========================================================
DEFAULT_DATA = {
    "kode_provinsi": "33",
    "nama_provinsi": "JAWA TENGAH",

    "kode_kabupaten_kota": "3374",
    "nama_kabupaten_kota": "KOTA SEMARANG",

    "kode_kecamatan": "000",
    "nama_kecamatan": "KECAMATAN",

    "kode_kelurahan_desa": "0000",
    "nama_kelurahan_desa": "KELURAHAN/DESA",

    # DATA PEMOHON
    "nama_lengkap": "-",
    "nik": "-",
    "nomor_kk_semula": "-",
    "alamat": "-",
    "rt": "000",
    "rw": "000",

    "desa_kelurahan": "-",
    "kecamatan": "-",
    "kabupaten_kota": "-",
    "provinsi": "-",

    "nomor_telepon": "-",
    "kode_pos": "00000",

    "alasan_permohonan": "",
    "jumlah_anggota_keluarga": "0",
    "anggota_keluarga": [],

    # TTD
    "tanggal": "",
    "tanggal_ttd": "",
    "tempat_ttd": "SRAGEN",

    "nama_kades": "KEPALA DESA",
    "nip_kades": "-",

    "qr_ttd": "",
}


# =========================================================
# MAPPING KODE SHDK
# =========================================================
SHDK_MAP = {
    "kepala-keluarga": "01",
    "suami": "02",
    "istri": "03",
    "anak": "04",
    "menantu": "05",
    "cucu": "06",
    "orang-tua": "07",
    "mertua": "08",
    "famili-lain": "09",
    "pembantu": "10",
    "lainnya": "11",
}


# =========================================================
# HELPER DASAR
# =========================================================
def normalize_text(value):
    if value is None:
        return ""
    return str(value).strip()


def safe_default(value, default_value=""):
    value = normalize_text(value)
    return value if value else default_value


def normalize_date(value):
    value = normalize_text(value)
    if not value:
        return ""

    for fmt in ("%Y-%m-%d", "%d-%m-%Y", "%Y/%m/%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(value, fmt).strftime("%d-%m-%Y")
        except ValueError:
            continue

    return value


def format_tanggal_ttd(tempat, tanggal):
    tempat = normalize_text(tempat)
    tanggal = normalize_date(tanggal)

    if tempat and tanggal:
        return f"{tempat}, {tanggal}"
    if tanggal:
        return tanggal
    if tempat:
        return tempat
    return ""


def get_real_cell_for_merged(ws, cell_ref):
    for merged_range in ws.merged_cells.ranges:
        if cell_ref in merged_range:
            return merged_range.start_cell.coordinate
    return cell_ref


def set_value(ws, cell_ref, value, align=None):
    if value is None or value == "":
        return
    real_cell = get_real_cell_for_merged(ws, cell_ref)
    ws[real_cell] = value
    if align:
        ws[real_cell].alignment = copy(align)


def check_cell(ws, cell_ref):
    real_cell = get_real_cell_for_merged(ws, cell_ref)
    ws[real_cell] = CHECK
    ws[real_cell].alignment = copy(CENTER)
    ws[real_cell].font = copy(FONT_CHECK)


def get_shdk_code(status_value):
    if not status_value:
        return ""
    return SHDK_MAP.get(normalize_text(status_value).lower(), "11")


def create_qr_image_file(text):
    if not text:
        return None

    temp_file = tempfile.NamedTemporaryFile(delete=False, suffix=".png")
    temp_file.close()

    qr = qrcode.QRCode(
        version=2,
        error_correction=qrcode.constants.ERROR_CORRECT_M,
        box_size=8,
        border=1,
    )
    qr.add_data(text)
    qr.make(fit=True)

    img = qr.make_image(fill_color="black", back_color="white")
    img.save(temp_file.name)
    return temp_file.name


def safe_remove_file(path):
    try:
        if path and os.path.exists(path):
            os.remove(path)
    except Exception:
        pass


def insert_qr_center_range(ws, image_path, start_cell, end_cell, width=65, height=65):
    if not image_path or not os.path.exists(image_path):
        return

    img = XLImage(image_path)
    img.width = width
    img.height = height

    start_col_letter, start_row = coordinate_from_string(start_cell)
    end_col_letter, end_row = coordinate_from_string(end_cell)

    start_col = column_index_from_string(start_col_letter)
    end_col = column_index_from_string(end_col_letter)

    total_width = 0
    for col in range(start_col, end_col + 1):
        col_letter = get_column_letter(col)
        total_width += int((ws.column_dimensions[col_letter].width or 8.43) * 7)

    total_height = 0
    for row in range(start_row, end_row + 1):
        total_height += int((ws.row_dimensions[row].height or 15) * 1.33)

    offset_x = max((total_width - width) // 2, 0)
    offset_y = max((total_height - height) // 2, 0)

    marker = AnchorMarker(
        col=start_col - 1,
        colOff=offset_x * 9525,
        row=start_row - 1,
        rowOff=offset_y * 9525,
    )

    size = XDRPositiveSize2D(
        cx=width * 9525,
        cy=height * 9525,
    )

    img.anchor = OneCellAnchor(_from=marker, ext=size)
    ws.add_image(img)


# =========================================================
# TRANSFORM RAW DATA
# =========================================================
def transform_raw_to_data(raw_data):
    data = dict(DEFAULT_DATA)
    data.update(raw_data or {})

    # =========================
    # WILAYAH (FALLBACK AMAN)
    # =========================
    data["kode_provinsi"] = safe_default(data.get("kode_provinsi"), DEFAULT_DATA["kode_provinsi"])
    data["nama_provinsi"] = safe_default(data.get("nama_provinsi"), DEFAULT_DATA["nama_provinsi"])

    data["kode_kabupaten_kota"] = safe_default(data.get("kode_kabupaten_kota"), DEFAULT_DATA["kode_kabupaten_kota"])
    data["nama_kabupaten_kota"] = safe_default(data.get("nama_kabupaten_kota"), DEFAULT_DATA["nama_kabupaten_kota"])

    data["kode_kecamatan"] = safe_default(data.get("kode_kecamatan"), DEFAULT_DATA["kode_kecamatan"])
    data["nama_kecamatan"] = safe_default(data.get("nama_kecamatan"), DEFAULT_DATA["nama_kecamatan"])

    data["kode_kelurahan_desa"] = safe_default(data.get("kode_kelurahan_desa"), DEFAULT_DATA["kode_kelurahan_desa"])
    data["nama_kelurahan_desa"] = safe_default(data.get("nama_kelurahan_desa"), DEFAULT_DATA["nama_kelurahan_desa"])

    # =========================
    # DATA PEMOHON
    # =========================
    data["nama_lengkap"] = safe_default(data.get("nama_lengkap"), "-")
    data["nik"] = safe_default(data.get("nik"), "-")
    data["nomor_kk_semula"] = safe_default(data.get("nomor_kk_semula"), "-")
    data["alamat"] = safe_default(data.get("alamat"), "-")

    data["rt"] = safe_default(data.get("rt"), "000")
    data["rw"] = safe_default(data.get("rw"), "000")

    data["kode_pos"] = safe_default(data.get("kode_pos"), "00000")
    data["nomor_telepon"] = safe_default(data.get("nomor_telepon"), "-")

    data["jumlah_anggota_keluarga"] = data.get("jumlah_anggota_keluarga") or "0"

    # =========================
    # KADES (fallback dari B01 style)
    # =========================
    nama_kades_raw = raw_data.get("nama_kades")
    if not normalize_text(nama_kades_raw):
        nama_kades_raw = raw_data.get("nama_kepala_desa")

    nip_kades_raw = raw_data.get("nip_kades")
    if not normalize_text(nip_kades_raw):
        nip_kades_raw = raw_data.get("nip_kepala_desa")

    data["nama_kades"] = safe_default(nama_kades_raw, DEFAULT_DATA["nama_kades"])
    data["nip_kades"] = safe_default(nip_kades_raw, DEFAULT_DATA["nip_kades"])

    # =========================
    # TANGGAL AUTO
    # =========================
    tanggal_display = safe_default(data.get("tanggal"))
    if not tanggal_display:
        tanggal_display = format_tanggal_ttd(
            data.get("tempat_ttd", DEFAULT_DATA["tempat_ttd"]),
            data.get("tanggal_ttd", ""),
        )
    data["tanggal"] = tanggal_display

    # =========================
    # QR
    # =========================
    qr_text = safe_default(data.get("qr_ttd"))
    if not qr_text:
        qr_text = f"Ditandatangani secara elektronik oleh {data['nama_kades']}"
        if data["nip_kades"] and data["nip_kades"] != "-":
            qr_text += f", NIP {data['nip_kades']}"

    data["qr_text"] = qr_text

    # =========================
    # ANGGOTA
    # =========================
    anggota_raw = raw_data.get("anggota_keluarga", None)

    if not isinstance(anggota_raw, list):
        anggota_raw = []

    anggota_normalized = []
    for item in anggota_raw:
        if not isinstance(item, dict):
            continue

        anggota_normalized.append({
            "nik": item.get("nik") or item.get("anggota_nik", ""),
            "nama": item.get("nama") or item.get("anggota_nama_lengkap", ""),
            "status": item.get("status") or item.get("anggota_status_hubungan_keluarga", ""),
        })

    data["anggota_keluarga"] = anggota_normalized

    print("RAW anggota_keluarga:", raw_data.get("anggota_keluarga"))
    print("RAW anggota_keluraga:", raw_data.get("anggota_keluraga"))
    print("NORMALIZED anggota_keluarga:", anggota_normalized)

    return data


# =========================================================
# FILL HEADER
# =========================================================
def fill_header(ws, data):
    # PROVINSI
    set_value(ws, "N10", data["kode_provinsi"], CENTER)
    set_value(ws, "S10", data["nama_provinsi"], LEFT)

    # KABUPATEN / KOTA
    set_value(ws, "N11", data["kode_kabupaten_kota"], CENTER)
    set_value(ws, "S11", data["nama_kabupaten_kota"], LEFT)

    # KECAMATAN
    set_value(ws, "N12", data["kode_kecamatan"], CENTER)
    set_value(ws, "S12", data["nama_kecamatan"], LEFT)

    # KELURAHAN / DESA
    set_value(ws, "N13", data["kode_kelurahan_desa"], CENTER)
    set_value(ws, "S13", data["nama_kelurahan_desa"], LEFT)


# =========================================================
# FILL PEMOHON
# =========================================================
def fill_pemohon(ws, data):
    set_value(ws, "J17", data["nama_lengkap"], LEFT)
    set_value(ws, "J19", data["nik"], LEFT)
    set_value(ws, "J21", data["nomor_kk_semula"], LEFT)
    set_value(ws, "J23", data["alamat"], LEFT)

    set_value(ws, "AJ23", data["rt"], CENTER)
    set_value(ws, "AR23", data["rw"], CENTER)

    set_value(ws, "P25", data["desa_kelurahan"], LEFT)
    set_value(ws, "AK25", data["kecamatan"], LEFT)
    set_value(ws, "P27", data["kabupaten_kota"], LEFT)
    set_value(ws, "AK27", data["provinsi"], LEFT)
    set_value(ws, "P29", data["kode_pos"], CENTER)
    set_value(ws, "AK29", data["nomor_telepon"], LEFT)


# =========================================================
# CHECKBOX ALASAN
# =========================================================
def fill_alasan(ws, data):
    alasan = data["alasan_permohonan"]
    if alasan == "rumah-tangga-baru":
        check_cell(ws, "J31")
    elif alasan == "kk-hilang-rusak":
        check_cell(ws, "J33")
    else:
        check_cell(ws, "X31")


# =========================================================
# TABEL ANGGOTA
# =========================================================
def fill_anggota(ws, data):
    start_row = 42
    row_gap = 2
    anggota_list = data.get("anggota_keluarga", [])

    for i, anggota in enumerate(anggota_list):
        row = start_row + (i * row_gap)

        set_value(ws, f"C{row}", i + 1, CENTER)
        set_value(ws, f"F{row}", anggota.get("nik", ""), LEFT)
        set_value(ws, f"W{row}", anggota.get("nama", ""), LEFT)
        set_value(ws, f"AV{row}", get_shdk_code(anggota.get("status", "")), CENTER)


# =========================================================
# FOOTER / TANDA TANGAN KEPALA DESA
# =========================================================
def fill_footer(ws, data):
    set_value(ws, "M35", data["jumlah_anggota_keluarga"], CENTER)
    set_value(ws, "AJ62", data["tanggal"], CENTER)
    set_value(ws, "AJ69", data["nama_lengkap"], CENTER)
    set_value(ws, "U71", data["nama_kades"], CENTER)
    set_value(ws, "U72", data["nip_kades"], CENTER)

    qr_path = create_qr_image_file(data.get("qr_text", ""))
    if qr_path:
        insert_qr_center_range(ws, qr_path, "T67", "T70", width=65, height=65)

    return qr_path


# =========================================================
# MAIN GENERATOR
# =========================================================
def generate_b03_excel(template_path, raw_data):
    data = transform_raw_to_data(raw_data)

    wb = load_workbook(template_path)
    ws = wb[SHEET_NAME]

    fill_header(ws, data)
    fill_pemohon(ws, data)
    fill_alasan(ws, data)
    fill_anggota(ws, data)

    qr_path = fill_footer(ws, data)

    output = io.BytesIO()
    try:
        wb.save(output)
        output.seek(0)
        return output.getvalue()
    finally:
        safe_remove_file(qr_path)